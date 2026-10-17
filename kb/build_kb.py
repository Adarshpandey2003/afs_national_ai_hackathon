"""Build the AFS knowledge base from the supplied hackathon materials (kb/raw/).

Usage (from the repo root, with the backend venv):
    python kb/build_kb.py              # parse everything, OCR image pages (cached in kb/ocr_cache/)
    python kb/build_kb.py --skip-ocr   # parse only; image-only pages keep their sparse PDF text

Outputs (kb/):
    sources.json      source metadata incl. evidence level and scope note
    pages.jsonl       one record per source page (what /kb/{source}/pages/{page} serves)
    chunks.jsonl      retrieval chunks (page-bounded, reference lists excluded)
    methods.json      S01 Table 1 (from the xlsx) + method overview (from the reference docx)
    fibers.json       fiber families with transmission ranges in µm (S02 p.6)
    cases.json        D01–D08 development and H01–H06 evaluation enquiries
    reference/        challenge brief, answer guidelines and the worked D01/D02 examples (Markdown)
"""
import argparse
import json
import re
import sys
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import docx
import openpyxl
import pymupdf
import yaml

from ocr import OCR_MODEL, ocr_page

KB_DIR = Path(__file__).resolve().parent
RAW_DIR = KB_DIR / "raw"
REF_DIR = KB_DIR / "reference"

DOCX_FILE = "AFS_Examples_and_Reference.docx"
XLSX_FILE = "AFS_method_comparison_table.xlsx"
CASES_FILE = "customer_scenarios_for_final_evaluation.txt"
BRIEF_FILE = "AI-based molecular media analysis by process-spectroscopy in-line.pdf"

SOURCES_FILE = KB_DIR / "sources.yaml"

# Pages whose embedded images cover at least this share of the page get a figure transcription.
FIGURE_AREA_THRESHOLD = 0.10
TABLE1_PAGE = ("S01", 6)  # replaced by the clean xlsx transcription (the PDF text of this page is scrambled)
MAX_CHUNK_CHARS = 1400
MAX_TEXT_PAGE_CHARS = 3000  # text / Markdown documents are split into "pages" of about this size


# ---------------------------------------------------------------- helpers

LIGATURES = {"ﬁ": "fi", "ﬂ": "fl", "ﬀ": "ff", "ﬃ": "ffi", "ﬄ": "ffl"}


def clean_block(text: str) -> str:
    for lig, plain in LIGATURES.items():
        text = text.replace(lig, plain)
    text = text.replace("­", "")
    text = re.sub(r"-\n(?=\w)", "-", text)
    text = re.sub(r"\s*\n\s*", " ", text)
    return re.sub(r"[ \t]{2,}", " ", text).strip()


def to_um(value: float, unit: str) -> float:
    unit = unit.replace("um", "µm")
    if unit == "nm":
        return value / 1000
    if unit == "µm":
        return value
    return 10000 / value  # wavenumber in cm⁻¹


def parse_range_um(text: str):
    """'780–2500 nm' -> [0.78, 2.5]; '4000–400 cm⁻¹' -> [2.5, 25.0]; None if no range is given."""
    m = re.search(r"(\d[\d ]*\.?\d*)\s*[–-]\s*(\d[\d ]*\.?\d*)\s*(nm|µm|um|cm⁻¹|cm-1)", text)
    if not m:
        return None
    a, b = (float(x.replace(" ", "")) for x in m.group(1, 2))
    lo, hi = sorted((to_um(a, m.group(3)), to_um(b, m.group(3))))
    return [round(lo, 3), round(hi, 3)]


def split_long(text: str, max_chars: int) -> list[str]:
    sentences = re.split(r"(?<=[.!?])\s+", text)
    parts, cur = [], ""
    for s in sentences:
        if cur and len(cur) + len(s) + 1 > max_chars:
            parts.append(cur)
            cur = s
        else:
            cur = f"{cur} {s}".strip()
    if cur:
        parts.append(cur)
    return parts


def chunk_blocks(blocks: list[str], max_chars: int = MAX_CHUNK_CHARS) -> list[str]:
    """Group a page's blocks into chunks; a short trailing block is repeated as overlap."""
    chunks, cur = [], []
    for block in blocks:
        for piece in ([block] if len(block) <= max_chars else split_long(block, max_chars)):
            if cur and sum(len(x) for x in cur) + len(piece) > max_chars:
                chunks.append("\n\n".join(cur))
                cur = [cur[-1]] if len(cur[-1]) < 300 else []
            cur.append(piece)
    if cur:
        chunks.append("\n\n".join(cur))
    return chunks


def image_area_share(page) -> float:
    area = 0.0
    for img in page.get_images(full=True):
        try:
            box = page.get_image_bbox(img)
        except Exception:
            continue
        if box.is_valid and not box.is_infinite:
            area += box.width * box.height
    return area / (page.rect.width * page.rect.height)


# ---------------------------------------------------------------- docx

def docx_items(path: Path):
    """Yield (kind, style, payload) in document order; payload is text or a list of table rows."""
    d = docx.Document(path)
    for el in d.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = docx.text.paragraph.Paragraph(el, d)
            if p.text.strip():
                yield "p", p.style.name, p.text.strip()
        elif tag == "tbl":
            t = docx.table.Table(el, d)
            yield "tbl", None, [[c.text.strip() for c in row.cells] for row in t.rows]


def docx_sections(path: Path) -> dict[str, list]:
    """Split the docx into its numbered Heading 1 sections, keyed by number ('1'..'5')."""
    sections, current = {"0": []}, "0"
    for item in docx_items(path):
        kind, style, payload = item
        if kind == "p" and style == "Heading 1":
            current = payload.split(".", 1)[0].strip()
            sections[current] = []
        sections[current].append(item)
    return sections


def md_table(rows: list[list[str]]) -> str:
    rows = [[c.replace("\n", " ").replace("|", "/") for c in r] for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def to_markdown(items) -> str:
    out, number = [], 0
    for kind, style, payload in items:
        if kind == "tbl":
            out.append(md_table(payload))
            number = 0
            continue
        if style == "List Number":
            number += 1
            out.append(f"{number}. {payload}")
            continue
        number = 0
        if style == "List Bullet":
            out.append(f"- {payload}")
        elif style and style.startswith("Heading"):
            level = int(style.split()[-1])
            out.append(f"{'#' * level} {payload}")
        else:
            out.append(payload)
    text = "\n\n".join(out)
    # keep list items on consecutive lines
    return re.sub(r"\n\n(?=(?:\d+\.|-) )", "\n", text)


def parse_dev_cases(items) -> list[dict]:
    cases, current = [], None
    for kind, style, payload in items:
        if kind != "p":
            continue
        m = re.match(r"^(D\d\d)\s*·\s*(.+)$", payload)
        if style == "Heading 2" and m:
            current = {"case_id": m.group(1), "set": "dev", "title": m.group(2).strip(),
                       "text": "", "facts_stated": [], "dev_hint_sources": ""}
            cases.append(current)
        elif current:
            if payload.startswith("Facts stated"):
                facts = payload[len("Facts stated"):].strip()
                current["facts_stated"] = [f.strip() for f in facts.split(";") if f.strip()]
            elif payload.startswith("Useful sources"):
                current["dev_hint_sources"] = payload[len("Useful sources"):].strip()
            elif not current["text"]:
                current["text"] = payload
    return cases


def parse_eval_cases(path: Path) -> tuple[list[dict], str]:
    segments = re.split(r"^=+\s*$", path.read_text(encoding="utf-8"), flags=re.M)
    preamble = segments[0].strip()
    cases = []
    for i, seg in enumerate(segments):
        m = re.fullmatch(r"\s*(H\d\d)\s*\|\s*(.+?)\s*", seg)
        if m and i + 1 < len(segments):
            body = " ".join(line.strip() for line in segments[i + 1].strip().splitlines())
            cases.append({"case_id": m.group(1), "set": "eval", "title": m.group(2), "text": body})
    return cases, preamble


# ---------------------------------------------------------------- tables

def load_table1(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True)
    sheets = wb.worksheets
    rows = [[("" if v is None else str(v)) for v in r] for r in sheets[0].iter_rows(values_only=True)]
    header, body = rows[0], [r for r in rows[1:] if any(r)]
    notes = []
    if len(sheets) > 1:
        notes = [str(r[0]) for r in sheets[1].iter_rows(values_only=True) if r and r[0]]
    return {"citation": "S01 Table 1 (p.6)", "columns": header,
            "rows": [dict(zip(header, r)) for r in body], "notes": notes}


def table1_page_text(table1: dict) -> str:
    lines = ["Table 1. Comparative characteristics of spectral methods and measurement techniques.",
             "(Clean transcription from AFS_method_comparison_table.xlsx; the PDF text of this page is scrambled.)",
             ""]
    for r in table1["rows"]:
        lines.append(f"- {r['method']} | {r['measurement_technique']} | phenomenon: {r['physical_phenomenon']} | "
                     f"sample: {r['sample_type']} | optical path or depth: {r['optical_path_or_depth']} | "
                     f"sensitivity: {r['sensitivity']} | selectivity: {r['selectivity']} | cost: {r['cost']}")
    lines += ["", "Notes: sensitivity, selectivity and cost are the source's qualitative ratings "
                  "(low / average / high), not measured values."]
    return "\n".join(lines)


def table1_chunk_blocks(table1: dict) -> list[str]:
    """One block per method, each repeating the table title so a chunk is self-explanatory."""
    by_method = {}
    for line in table1_page_text(table1).splitlines():
        if line.startswith("- "):
            by_method.setdefault(line[2:].split(" | ", 1)[0], []).append(line)
    return [f"S01 Table 1 (comparative characteristics of spectral methods), {method}:\n" + "\n".join(rows)
            for method, rows in by_method.items()]


def method_overview(items) -> dict:
    tables = [p for k, _, p in items if k == "tbl"]
    paragraphs = [p for k, s, p in items if k == "p" and s == "Normal"]
    overview_rows, fiber_rows = tables[0], tables[1]
    methods = []
    for method, rng, samples, limits, read in overview_rows[1:]:
        methods.append({"method": method, "nominal_range": rng, "range_um": parse_range_um(rng),
                        "typical_samples_and_geometry": samples, "main_limitations": limits,
                        "read": read})
    fibers = []
    for fiber_type, rng in fiber_rows[1:]:
        lo_hi = parse_range_um(rng)
        fibers.append({"fiber_type": fiber_type, "range_text": rng,
                       "min_um": lo_hi[0] if lo_hi else None, "max_um": lo_hi[1] if lo_hi else None})
    return {
        "methods": {"note": paragraphs[0] if paragraphs else "", "rows": methods},
        "fibers": {"citation": "S02 p.6",
                   "note": next((p for p in paragraphs if p.startswith("The fiber determines")), ""),
                   "unit_reminder": next((p for p in paragraphs if p.startswith("Unit reminder")), ""),
                   "fibers": fibers},
    }


# ---------------------------------------------------------------- pdf sources

def extract_pdf_pages(src: dict) -> list[dict]:
    doc = pymupdf.open(RAW_DIR / src["file"])
    pages, in_references = [], False
    for i, page in enumerate(doc, start=1):
        blocks = [clean_block(b[4]) for b in page.get_text("blocks") if b[6] == 0]
        blocks = [b for b in blocks if b and not re.fullmatch(r"\d{1,3}", b)]
        body = []
        for b in blocks:
            if re.fullmatch(r"(■\s*)?references?", b, flags=re.I):
                in_references = True
            if not in_references:
                body.append(b)
        pages.append({"page": i, "blocks": blocks, "body_blocks": body,
                      "image_share": image_area_share(page)})
    return pages


def load_sources() -> list[dict]:
    return yaml.safe_load(SOURCES_FILE.read_text(encoding="utf-8"))


def ocr_jobs(src: dict, pages: list[dict]) -> list[tuple]:
    """(page, mode) pairs to transcribe, from the source's `ocr` setting."""
    mode = src.get("ocr", "none")
    if mode == "slides":
        return [(p["page"], "slide") for p in pages]
    if mode == "figures":
        return [(p["page"], "figures") for p in pages
                if p["image_share"] >= FIGURE_AREA_THRESHOLD and (src["source_id"], p["page"]) != TABLE1_PAGE]
    return []


def run_ocr(src: dict, jobs: list[tuple]) -> dict[int, str]:
    def run(job):
        page_no, mode = job
        return page_no, ocr_page(RAW_DIR / src["file"], src["source_id"], page_no, mode)

    with ThreadPoolExecutor(max_workers=4) as pool:
        return dict(pool.map(run, jobs))


def text_pages(text: str, max_chars: int = MAX_TEXT_PAGE_CHARS) -> list[list[str]]:
    """Group the paragraphs of a text / Markdown document into pages of at most about max_chars."""
    pages, cur = [], []
    for para in (p.strip() for p in re.split(r"\n\s*\n", text)):
        if not para:
            continue
        if cur and sum(map(len, cur)) + len(para) > max_chars:
            pages.append(cur)
            cur = []
        cur.append(para)
    return pages + [cur] if cur else pages


def source_records(src: dict, table1: dict | None = None, skip_ocr: bool = False) -> tuple[list[dict], list[dict]]:
    """Page records and retrieval chunks for one source: a PDF, or a text / Markdown document."""
    sid, pages, chunks = src["source_id"], [], []

    def add_page(page_no, text, origin, chunk_source_blocks):
        label = f"{sid} p.{page_no}"
        pages.append({"source": sid, "page": page_no, "label": label, "title": src["title"],
                      "evidence_level": src["evidence_level"], "origin": origin, "text": text})
        for n, chunk in enumerate(chunk_blocks(chunk_source_blocks)):
            chunks.append({"chunk_id": f"{sid}-p{page_no:02d}-c{n}", "source": sid, "page": page_no,
                           "label": label, "evidence_level": src["evidence_level"], "origin": origin, "text": chunk})

    path = RAW_DIR / src["file"]
    if path.suffix.lower() != ".pdf":
        for n, paragraphs in enumerate(text_pages(path.read_text(encoding="utf-8")), start=1):
            add_page(n, "\n\n".join(paragraphs), "text_file", paragraphs)
        return pages, chunks

    extracted = extract_pdf_pages(src)
    ocr = {} if skip_ocr else run_ocr(src, ocr_jobs(src, extracted))
    for p in extracted:
        page_no = p["page"]
        if (sid, page_no) == TABLE1_PAGE and table1:
            add_page(page_no, table1_page_text(table1), "table_transcription", table1_chunk_blocks(table1))
            continue
        pdf_text = "\n\n".join(p["blocks"])
        vision = ocr.get(page_no, "")
        if src.get("ocr") == "slides" and vision:
            text = (f"[Slide text layer]\n{pdf_text}\n\n" if pdf_text else "") + \
                   f"[Vision transcription ({OCR_MODEL})]\n{vision}"
            add_page(page_no, text, "vision_ocr", [b for b in text.split("\n\n") if b.strip()])
        elif vision and "NO_INFORMATIVE_FIGURE" not in vision:
            figure = f"[Figure content, vision transcription ({OCR_MODEL})]\n{vision}"
            add_page(page_no, f"{pdf_text}\n\n{figure}", "pdf_text+figure_ocr",
                     p["body_blocks"] + [b for b in figure.split("\n\n") if b.strip()])
        else:
            add_page(page_no, pdf_text, "pdf_text", p["body_blocks"])
    return pages, chunks


# ---------------------------------------------------------------- main

def main(skip_ocr: bool) -> None:
    REF_DIR.mkdir(exist_ok=True)
    sections = docx_sections(RAW_DIR / DOCX_FILE)

    # source descriptions from the reference doc, section 4 table
    src_table = next(p for k, _, p in sections["4"] if k == "tbl")
    descriptions = {row[0]: row[2] for row in src_table[1:]}

    table1 = load_table1(RAW_DIR / XLSX_FILE)
    overview = method_overview(sections["5"])

    sources = load_sources()
    pages_out, chunks_out = [], []
    for src in sources:
        pages, chunks = source_records(src, table1, skip_ocr)
        pages_out += pages
        chunks_out += chunks

    page_counts = {}
    for rec in pages_out:
        page_counts[rec["source"]] = max(page_counts.get(rec["source"], 0), rec["page"])
    sources_out = [{**s, "pages": page_counts.get(s["source_id"], 0),
                    "description": s.get("description") or descriptions.get(s["source_id"], "")} for s in sources]

    dev_cases = parse_dev_cases(sections["1"])
    eval_cases, eval_preamble = parse_eval_cases(RAW_DIR / CASES_FILE)
    scenario_notes = [p for k, s, p in sections["1"] if k == "p" and s == "List Bullet"]

    brief = pymupdf.open(RAW_DIR / BRIEF_FILE)
    (REF_DIR / "challenge_brief.md").write_text(
        "# Challenge brief (extracted text)\n\n" + "\n\n".join(pg.get_text() for pg in brief), encoding="utf-8")
    (REF_DIR / "guidelines.md").write_text(
        "# Answer guidelines (AFS Examples and Reference, sections 1, 3 and 4)\n\n"
        "## Notes that apply to all scenarios\n\n" + "\n".join(f"- {n}" for n in scenario_notes) + "\n\n"
        + to_markdown(sections["3"]) + "\n\n" + to_markdown(sections["4"]) + "\n", encoding="utf-8")
    (REF_DIR / "worked_examples.md").write_text(to_markdown(sections["2"]) + "\n", encoding="utf-8")

    def dump_json(name, obj):
        (KB_DIR / name).write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")

    def dump_jsonl(name, records):
        with open(KB_DIR / name, "w", encoding="utf-8") as f:
            for r in records:
                f.write(json.dumps(r, ensure_ascii=False) + "\n")

    dump_json("sources.json", sources_out)
    dump_jsonl("pages.jsonl", pages_out)
    dump_jsonl("chunks.jsonl", chunks_out)
    dump_json("methods.json", {"table1": table1, "method_overview": overview["methods"]})
    dump_json("fibers.json", overview["fibers"])
    dump_json("cases.json", {"notes": eval_preamble, "cases": dev_cases + eval_cases})

    print(f"sources: {len(sources_out)}  pages: {len(pages_out)}  chunks: {len(chunks_out)}  "
          f"cases: {len(dev_cases)} dev + {len(eval_cases)} eval")
    origins = {}
    for rec in pages_out:
        origins[rec["origin"]] = origins.get(rec["origin"], 0) + 1
    print("page origins:", origins)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--skip-ocr", action="store_true", help="do not call the vision model")
    main(parser.parse_args().skip_ocr)
